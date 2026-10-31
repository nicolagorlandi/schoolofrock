import json
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth.models import AnonymousUser
from django.contrib.auth.models import Permission, User
from django.core.cache import cache
from django.test import RequestFactory, SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from documenti.exports import esporta_tabelle_excel
from documenti.forms import creaDettaglioOrdineForm, creaOrdineForm
from documenti.models import (
    Azienda,
    Articoli,
    Causdocu,
    Codiciva,
    MenuItem,
    Nominativi,
    Ord_Mast,
    Ord_dett,
    PAG_DETT,
    PAGAMENTI,
)
from documenti.views import nominativi, nominativi_list_url
from documenti.services import aggiorna_riga_cassa
from PaNet1Project import settings


class AziendaCacheTests(SimpleTestCase):
    cache_key = "documenti_azienda"

    def setUp(self):
        cache.delete(self.cache_key)

    def tearDown(self):
        cache.delete(self.cache_key)

    def test_get_azienda_reads_from_database_once_and_caches_without_expiry(self):
        azienda = Azienda(NOME="Test")
        with patch.object(Azienda.objects, "order_by") as order_by:
            order_by.return_value.first.return_value = azienda

            self.assertEqual(Azienda.get_azienda(), azienda)
            self.assertEqual(Azienda.get_azienda(), azienda)

        order_by.assert_called_once_with("NOME")

    def test_save_invalidates_cached_azienda(self):
        azienda = Azienda(NOME="Test")
        cache.set(self.cache_key, azienda, timeout=None)

        with patch("django.db.models.Model.save"), patch(
            "documenti.models.transaction.on_commit",
            side_effect=lambda callback: callback(),
        ):
            azienda.save()

        self.assertIsNone(cache.get(self.cache_key))


class DocumentoFormTests(TestCase):
    def test_edit_preserves_document_type_when_hidden_field_is_not_posted(self):
        operator = User.objects.create_user(username="documento-form-test")
        document = Ord_Mast.objects.create(
            DOCODUTE=operator,
            DOTIPDOC="FV",
            DO_CASSA=50,
            slug="documento-form-test",
        )

        form = creaOrdineForm(
            data={},
            instance=document,
            TipoDocumentoSelezionato="FV",
        )

        self.assertIn(
            "DOTIPDOC",
            [field.name for field in form.hidden_fields()],
        )
        self.assertTrue(form.fields["DOTIPDOC"].disabled)
        self.assertTrue(form.is_valid(), form.errors)
        updated_document = form.save()

        self.assertEqual(updated_document.DOTIPDOC, "FV")
        self.assertEqual(updated_document.DO_CASSA, Decimal("50.00"))

    def setUp(self):
        self.iva_articolo = Codiciva.objects.create(
            CICODICE="Q22T1",
            CIDESCRI="IVA 22%",
            slug="q22t1",
        )
        self.iva_alternativa = Codiciva.objects.create(
            CICODICE="Q10T1",
            CIDESCRI="IVA 10%",
            slug="q10t1",
        )
        self.articolo = Articoli.objects.create(
            ARCODICE="ART-IVA",
            ARDESCRI="Articolo IVA",
            ARCODIVA=self.iva_articolo,
            slug="art-iva",
        )

    def test_detail_defaults_to_article_vat_and_accepts_a_manual_choice(self):
        form_data = {
            "DOCODART": self.articolo.pk,
            "DOQTAUM1": "1",
        }
        form = creaDettaglioOrdineForm(data=form_data)

        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["DOCODIVA"], self.iva_articolo)
        self.assertEqual(
            [field.name for field in form.visible_fields()][-1],
            "DOCODIVA",
        )

        form_data["DOCODIVA"] = self.iva_alternativa.pk
        form = creaDettaglioOrdineForm(data=form_data)

        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["DOCODIVA"], self.iva_alternativa)

    def test_article_lookup_returns_its_vat_code(self):
        response = self.client.get(
            reverse("leggiPrezzoDescrizioneArticolo"),
            {
                "articoloId": self.articolo.pk,
                "name": "ord_dett_set-0-DOCODART",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["ARCODIVA"], self.iva_articolo.pk)

    def test_document_totals_use_the_selected_line_vat(self):
        operator = User.objects.create_user(username="vat-total-test")
        causale = Causdocu.objects.create(
            CDCODICE="QVAT1",
            CDDESCRI="Causale IVA",
            CDTIPDOC="FV",
            CDQUALIT=0,
            slug="qvat1",
        )
        document = Ord_Mast.objects.create(
            DOCODUTE=operator,
            DOCODCAU=causale,
            DOTIPDOC="FV",
            slug="vat-total-test",
        )
        Ord_dett.objects.create(
            DOCODICE=document,
            DOCODART=self.articolo,
            DOCODIVA=self.iva_alternativa,
            DOQTAUM1=1,
            DOPREZZO=100,
        )

        riepiloghi, _ = document.totale_fattura()

        self.assertEqual(len(riepiloghi), 1)
        self.assertEqual(riepiloghi[0].codice_iva, self.iva_alternativa.pk)
        self.assertEqual(riepiloghi[0].percentuale_iva, self.iva_alternativa.CIPERIVA)


class RigaCassaAutomaticaTests(TestCase):
    def setUp(self):
        self.iva = Codiciva.objects.create(
            CICODICE="QCAS1",
            CIDESCRI="IVA 22%",
            CIPERIVA=22,
            slug="qcas1",
        )
        self.articolo_cassa = Articoli.objects.create(
            ARCODICE="ART-CASSA",
            ARDESCRI="Cassa previdenziale",
            ARCODIVA=self.iva,
            slug="art-cassa",
        )
        self.azienda = Azienda.objects.order_by("NOME").first()
        self.azienda.AZCODCAS = self.articolo_cassa
        self.azienda.save(update_fields=["AZCODCAS"])
        cache.delete("documenti_azienda")
        self.intestatario = Nominativi.objects.create(
            NOCODICE="CASSA-CLIENTE",
            NORAGSOC="Cliente cassa",
            NOCASSA=4,
        )
        self.causale = Causdocu.objects.create(
            CDCODICE="CAS01",
            CDDESCRI="Fattura",
            CDTIPDOC="FV",
            CDQUALIT=0,
            slug="cas01",
        )
        self.documento = Ord_Mast.objects.create(
            DOCODUTE=User.objects.create_user(username="cassa-test"),
            DOCODCAU=self.causale,
            DOCODINT=self.intestatario,
            DOTIPDOC="FV",
            DO_CASSA=50,
            slug="cassa-documento",
        )
        Ord_dett.objects.create(
            DOCODICE=self.documento,
            DONUMRIG=1,
            DOCODART=self.articolo_cassa,
            DOQTAUM1=1,
            DOPREZZO=200,
            DORIGCASSA=False,
        )

    def tearDown(self):
        cache.delete("documenti_azienda")

    def test_saves_and_recalculates_single_automatic_cassa_line(self):
        aggiorna_riga_cassa(self.documento)
        riga = Ord_dett.objects.get(DORIGCASSA=True)

        self.assertEqual(riga.DOCODART, self.articolo_cassa)
        self.assertEqual(riga.DOCODIVA, self.iva)
        self.assertEqual(riga.DOPREZZO, Decimal("4.00"))

        self.documento.DO_CASSA = Decimal("100")
        aggiorna_riga_cassa(self.documento)
        righe_cassa = Ord_dett.objects.filter(DORIGCASSA=True)

        self.assertEqual(righe_cassa.count(), 1)
        self.assertEqual(righe_cassa.get().DOPREZZO, Decimal("8.00"))

    def test_removes_automatic_line_when_percentage_is_disabled(self):
        aggiorna_riga_cassa(self.documento)
        self.documento.DO_CASSA = Decimal("0")

        aggiorna_riga_cassa(self.documento)

        self.assertFalse(Ord_dett.objects.filter(DORIGCASSA=True).exists())

    def test_calculates_cassa_on_the_vat_excluded_taxable_base(self):
        self.documento.DOLORIVA = "S"

        aggiorna_riga_cassa(self.documento)

        self.assertEqual(
            Ord_dett.objects.get(DORIGCASSA=True).DOPREZZO,
            Decimal("3.28"),
        )

    def test_only_selected_document_types_receive_automatic_line(self):
        self.documento.DOTIPDOC = "OR"

        aggiorna_riga_cassa(self.documento)

        self.assertFalse(Ord_dett.objects.filter(DORIGCASSA=True).exists())


class DocumentoPrintTests(TestCase):
    def setUp(self):
        cache.delete("documenti_azienda")
        self.operatore = User.objects.create_user(username="documento-print-test")
        self.azienda = Azienda.objects.order_by("NOME").first()
        self.azienda.NOTE = "Emittente riga uno\nEmittente riga due"
        self.azienda.save(update_fields=["NOTE"])
        cache.delete("documenti_azienda")

    def tearDown(self):
        cache.delete("documenti_azienda")

    def _render_document(self, ciclo):
        causale = Causdocu.objects.create(
            CDCODICE=f"PR{ciclo}",
            CDDESCRI="Ricevuta fiscale",
            CD_CICLO=ciclo,
            CDQUALIT=0,
            slug=f"print-{ciclo.lower()}",
        )
        documento = Ord_Mast.objects.create(
            DOCODUTE=self.operatore,
            DOCODCAU=causale,
            DONUMORD=42,
            slug=f"print-{ciclo.lower()}-doc",
        )

        with patch("documenti.views.pdfkit.configuration"), patch(
            "documenti.views.pdfkit.from_string",
            return_value=b"%PDF",
        ) as pdf_generator:
            response = self.client.get(
                reverse("stampaDocumento", kwargs={"pk": documento.pk})
            )

        self.assertEqual(response.status_code, 200)
        return pdf_generator.call_args.args[0]

    def test_print_uses_cause_description_and_shows_issuer_for_active_cycle(self):
        html = self._render_document("V")

        self.assertIn("<h2>Ricevuta fiscale 42</h2>", html)
        self.assertIn("Emittente", html)
        self.assertIn("Emittente riga uno<br>Emittente riga due", html)

    def test_print_does_not_show_issuer_for_passive_cycle(self):
        html = self._render_document("A")

        self.assertIn("<h2>Ricevuta fiscale 42</h2>", html)
        self.assertNotIn("Emittente", html)


class EsportazioneExcelTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()

    def post(self, payload, user=None):
        payload.setdefault('page_name', 'elenco')
        request = self.factory.post(
            '/export/excel/',
            data=json.dumps(payload),
            content_type='application/json',
        )
        request.user = user or SimpleNamespace(is_authenticated=True)
        return esporta_tabelle_excel(request)

    def test_esporta_tabelle_su_fogli_separati_e_protegge_le_formule(self):
        response = self.post({
            'sheets': [
                {'name': 'Riepilogo', 'rows': [['Corso', 'Costo'], ['=1+1', '12,50']]},
                {'name': 'Riepilogo', 'rows': [['Alunno'], ['Mario Rossi']]},
            ]
        })

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Riepilogo', 'Riepilogo (2)'])
        self.assertEqual(workbook['Riepilogo']['A2'].value, '=1+1')
        self.assertEqual(workbook['Riepilogo']['A2'].data_type, 's')
        self.assertEqual(workbook['Riepilogo']['B2'].value, 12.5)
        self.assertEqual(workbook['Riepilogo (2)']['A2'].value, 'Mario Rossi')

    @patch('documenti.exports.timezone.localtime')
    def test_nome_file_usa_pagina_e_timestamp(self, localtime):
        localtime.return_value = datetime(2025, 1, 2, 3, 4, 5)
        response = self.post({
            'page_name': 'presenze-mensili',
            'sheets': [{'name': 'Riepilogo', 'rows': [['Corso'], ['Chitarra']]}],
        })

        self.assertEqual(
            response['Content-Disposition'],
            'attachment; filename="presenze-mensili_20250102_030405.xlsx"',
        )

    def test_rifiuta_nome_pagina_non_valido(self):
        response = self.post({
            'page_name': '../presenze',
            'sheets': [{'name': 'Elenco', 'rows': [['Corso'], ['Chitarra']]}],
        })

        self.assertEqual(response.status_code, 400)

    def test_richiede_autenticazione(self):
        response = self.post(
            {'sheets': [{'name': 'Elenco', 'rows': [['Colonna'], ['Dato']]}]},
            user=SimpleNamespace(is_authenticated=False),
        )

        self.assertEqual(response.status_code, 302)

    def test_rifiuta_payload_non_valido(self):
        response = self.post({'sheets': []})

        self.assertEqual(response.status_code, 400)


class NominativiListUrlTests(SimpleTestCase):
    def test_cliente_and_fornitore_types_resolve_to_existing_lists(self):
        self.assertEqual(nominativi_list_url(settings.CCATCLI), '/doc/clienti')
        self.assertEqual(nominativi_list_url(settings.CCATFOR), '/doc/fornitori')

    def test_unsupported_type_is_reported(self):
        with self.assertRaises(ValueError):
            nominativi_list_url('OTHER')


class NominativiSearchTests(TestCase):
    def test_search_filters_combine_across_requested_fields(self):
        match = Nominativi.objects.create(
            NOCODICE='ROSSI_MARIO',
            NOCODNOM=settings.CCATCLI,
            slug='rossi-mario',
            NOCODFIS='RSSMRA80A01H501U',
            NOPARIVA='01234567890',
            NO_NOME='Mario',
            NOCOGNOM='Rossi',
            NOTELEF2='3331234567',
            NO_CITTA='Roma',
        )
        Nominativi.objects.create(
            NOCODICE='BIANCHI_MARIO',
            NOCODNOM=settings.CCATCLI,
            slug='bianchi-mario',
            NOCODFIS='BNCMRA80A01H501U',
            NO_NOME='Mario',
            NOCOGNOM='Bianchi',
            NOTELEF1='3337654321',
            NO_CITTA='Milano',
        )

        request = RequestFactory().get(
            '/doc/clienti',
            {
                'nome': 'mar',
                'cognome': 'ross',
                'telefono': '333123',
                'citta': 'rom',
                'codice_fiscale': 'RSSMRA',
                'partita_iva': '012345',
            },
        )
        request.user = AnonymousUser()

        with patch('documenti.views.render') as render:
            nominativi(request, settings.CCATCLI)

        context = render.call_args.args[2]
        self.assertEqual(list(context['Nominativi']), [match])
        self.assertEqual(context['filtri']['citta'], 'rom')


class CausdocuManagementTests(TestCase):
    def setUp(self):
        user = User.objects.create_user(username='causdocu-manager')
        permission = Permission.objects.get(
            content_type__app_label='documenti',
            codename='Manager_Ordini',
        )
        user.user_permissions.add(permission)
        self.client.force_login(user)
        self.user = user

    def causale_form_data(self, **overrides):
        data = {
            'CDCODICE': 'TST01',
            'CDDESCRI': 'Causale test',
            'CD_CICLO': 'V',
            'CDTIPDOC': 'RF',
            'CDTIPINT': settings.CCATCLI,
            'CDCAUMAG': '',
            'CDNUMAUT': '1',
            'CDQUALIT': '0',
            'CDINTOBB': '1',
            'CDGESIVA': '1',
            'CDGESPAG': '1',
        }
        data.update(overrides)
        return data

    def pagamento_formset_data(self, rows=(), initial_forms=0):
        data = {
            'dettagli-TOTAL_FORMS': str(len(rows)),
            'dettagli-INITIAL_FORMS': str(initial_forms),
            'dettagli-MIN_NUM_FORMS': '0',
            'dettagli-MAX_NUM_FORMS': '1000',
        }
        for index, row in enumerate(rows):
            for field, value in row.items():
                data[f'dettagli-{index}-{field}'] = str(value)
        return data

    def test_menu_places_document_causals_under_archives_documents(self):
        voce = MenuItem.objects.get(
            label='Causali Documenti',
            url_name='causali-documenti-list',
        )

        self.assertEqual(voce.parent.label, 'Documenti')
        self.assertEqual(voce.parent.parent.label, 'Archivi')

    def test_menu_places_payment_methods_under_archives(self):
        voce = MenuItem.objects.get(
            label='Pagamenti',
            url_name='pagamenti-list',
        )

        self.assertEqual(voce.parent.label, 'Archivi')

    def test_navigation_has_a_mobile_menu_toggle(self):
        response = self.client.get(reverse('pagamenti-list'))

        self.assertContains(response, 'data-bs-target="#mainMenu"')
        self.assertContains(response, 'aria-controls="mainMenu"')
        self.assertContains(response, 'Apri il menu di navigazione')

    def test_create_update_and_list_payment_methods(self):
        form_response = self.client.get(reverse('pagamenti-create'))
        self.assertEqual(form_response.status_code, 200)
        self.assertContains(form_response, 'id_dettagli-TOTAL_FORMS')
        self.assertContains(form_response, 'Aggiungi rata')

        response = self.client.post(
            reverse('pagamenti-create'),
            {
                'PACODICE': 'BON',
                'PADESCRI': 'Bonifico bancario',
                'PADESEAS': 'Bank transfer',
                **self.pagamento_formset_data(),
            },
        )

        self.assertRedirects(response, reverse('pagamenti-list'))
        pagamento = PAGAMENTI.objects.get(pk='BON')
        self.assertEqual(pagamento.slug, 'bon')
        self.assertEqual(pagamento.PADESEAS, 'Bank transfer')

        response = self.client.post(
            reverse('pagamenti-update', args=[pagamento.pk]),
            {
                'PACODICE': 'BON',
                'PADESCRI': 'Bonifico aggiornato',
                'PADESEAS': 'Updated transfer',
                **self.pagamento_formset_data(),
            },
        )

        self.assertRedirects(response, reverse('pagamenti-list'))
        pagamento.refresh_from_db()
        self.assertEqual(pagamento.PADESCRI, 'Bonifico aggiornato')
        self.assertContains(
            self.client.get(reverse('pagamenti-list')),
            'Updated transfer',
        )

    def test_create_and_update_payment_installments(self):
        response = self.client.post(
            reverse('pagamenti-create'),
            {
                'PACODICE': 'RATE',
                'PADESCRI': 'Pagamento rateale',
                'PADESEAS': '',
                **self.pagamento_formset_data([{
                    'PDNUMSCA': 1,
                    'PDGIORNO': 'Data Doc.',
                    'PDINTERV': 1,
                    'PDPERCEN': '50.0000',
                    'PDPAGATA': 'N',
                    'PD_FISSO': 0,
                }]),
            },
        )

        self.assertRedirects(response, reverse('pagamenti-list'))
        pagamento = PAGAMENTI.objects.get(pk='RATE')
        rata = PAG_DETT.objects.get(PACODICE=pagamento, PDNUMSCA=1)
        self.assertEqual(rata.PDPERCEN, Decimal('50.0000'))

        response = self.client.post(
            reverse('pagamenti-update', args=[pagamento.pk]),
            {
                'PACODICE': 'RATE',
                'PADESCRI': 'Pagamento rateale aggiornato',
                'PADESEAS': '',
                **self.pagamento_formset_data(
                    [{
                        'id': rata.pk,
                        'PDNUMSCA': 1,
                        'PDGIORNO': 'Fine Mese',
                        'PDINTERV': 2,
                        'PDPERCEN': '100.0000',
                        'PDPAGATA': 'S',
                        'PD_FISSO': 10,
                    }],
                    initial_forms=1,
                ),
            },
        )

        self.assertRedirects(response, reverse('pagamenti-list'))
        rata.refresh_from_db()
        self.assertEqual(rata.PDGIORNO, 'Fine Mese')
        self.assertEqual(rata.PDPERCEN, Decimal('100.0000'))
        self.assertEqual(rata.PDPAGATA, 'S')
        self.assertContains(
            self.client.get(reverse('pagamenti-list')),
            'Fine Mese',
        )

    def test_delete_payment_installment_from_update_form(self):
        pagamento = PAGAMENTI.objects.create(
            PACODICE='DELRT',
            PADESCRI='Pagamento da aggiornare',
            slug='delrt',
        )
        rata = PAG_DETT.objects.create(
            PACODICE=pagamento,
            PDNUMSCA=1,
            PDPERCEN=50,
        )

        response = self.client.post(
            reverse('pagamenti-update', args=[pagamento.pk]),
            {
                'PACODICE': 'DELRT',
                'PADESCRI': 'Pagamento da aggiornare',
                'PADESEAS': '',
                **self.pagamento_formset_data(
                    [{
                        'id': rata.pk,
                        'PDNUMSCA': 1,
                        'DELETE': 'on',
                    }],
                    initial_forms=1,
                ),
            },
        )

        self.assertRedirects(response, reverse('pagamenti-list'))
        self.assertFalse(PAG_DETT.objects.filter(pk=rata.pk).exists())

    def test_cannot_delete_payment_method_used_by_a_document_causal(self):
        pagamento = PAGAMENTI.objects.create(
            PACODICE='USEDP',
            PADESCRI='Metodo utilizzato',
            slug='usedp',
        )
        causale = Causdocu.objects.create(
            CDCODICE='USE01',
            CDDESCRI='Causale pagamento',
            CDTIPDOC='RF',
            CDQUALIT=0,
            CDPAGDEF=pagamento,
            slug='use01',
        )

        response = self.client.post(
            reverse('pagamenti-delete', args=[pagamento.pk]),
            follow=True,
        )

        self.assertContains(response, 'è utilizzato da documenti')
        self.assertTrue(PAGAMENTI.objects.filter(pk=pagamento.pk).exists())
        self.assertTrue(Causdocu.objects.filter(pk=causale.pk).exists())

    def test_create_and_update_document_causal(self):
        pagamento = PAGAMENTI.objects.create(
            PACODICE='TESTP',
            PADESCRI='Pagamento predefinito test',
            slug='testp',
        )
        response = self.client.post(
            reverse('causali-documenti-create'),
            self.causale_form_data(CDPAGDEF=pagamento.pk),
        )

        self.assertRedirects(response, reverse('causali-documenti-list'))
        causale = Causdocu.objects.get(pk='TST01')
        self.assertEqual(causale.slug, 'tst01')
        self.assertEqual(causale.CDPAGDEF, pagamento)

        response = self.client.post(
            reverse('causali-documenti-update', args=[causale.pk]),
            self.causale_form_data(
                CDDESCRI='Descrizione aggiornata',
                CDPAGDEF=pagamento.pk,
            ),
        )

        self.assertRedirects(response, reverse('causali-documenti-list'))
        causale.refresh_from_db()
        self.assertEqual(causale.CDDESCRI, 'Descrizione aggiornata')

    def test_list_filters_document_causals_by_cycle_type_description_and_code(self):
        Causdocu.objects.create(
            CDCODICE='RF001',
            CDDESCRI='Ricevuta fiscale',
            CD_CICLO='V',
            CDTIPDOC='RF',
            CDQUALIT=0,
            slug='rf001',
        )
        Causdocu.objects.create(
            CDCODICE='FA001',
            CDDESCRI='Fattura acquisto',
            CD_CICLO='A',
            CDTIPDOC='FA',
            CDQUALIT=0,
            slug='fa001',
        )

        response = self.client.get(
            reverse('causali-documenti-list'),
            {
                'ciclo': 'V',
                'tipo_documento': 'rf',
                'descrizione': 'fisc',
                'codice': '001',
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            list(response.context['causali_documento'].values_list(
                'CDCODICE',
                flat=True,
            )),
            ['RF001'],
        )
        self.assertContains(response, 'value="fisc"')
        self.assertContains(response, 'value="001"')

    def test_selected_causal_returns_its_default_payment(self):
        pagamento = PAGAMENTI.objects.create(
            PACODICE='CAUPG',
            PADESCRI='Pagamento causale',
            slug='caupg',
        )
        causale = Causdocu.objects.create(
            CDCODICE='CAU01',
            CDDESCRI='Causale con pagamento',
            CDTIPDOC='RF',
            CDQUALIT=0,
            CDPAGDEF=pagamento,
            slug='cau01',
        )

        response = self.client.get(
            reverse('leggiCausale'),
            {'causaleId': causale.pk},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {'DOCODPAG': pagamento.pk})

        causale.CDPAGDEF = None
        causale.save(update_fields=['CDPAGDEF'])
        response = self.client.get(
            reverse('leggiCausale'),
            {'causaleId': causale.pk},
        )
        self.assertEqual(response.json(), {'DOCODPAG': None})

    def test_cannot_delete_document_causal_used_by_a_document(self):
        causale = Causdocu.objects.create(
            CDCODICE='TST02',
            CDDESCRI='Causale in uso',
            CDTIPDOC='RF',
            CDQUALIT=0,
            slug='tst02',
        )
        nominativo = Nominativi.objects.create(
            NOCODICE='TST-CLIENT',
            NORAGSOC='Cliente test',
            NOCODNOM=settings.CCATCLI,
        )
        documento = Ord_Mast.objects.create(
            DOCODCAU=causale,
            DOCODINT=nominativo,
            DOCODUTE=self.user,
            DOTIPDOC='RF',
            slug='documento-test-causale',
        )

        response = self.client.post(
            reverse('causali-documenti-delete', args=[causale.pk]),
            follow=True,
        )

        self.assertContains(response, 'è utilizzata da documenti')
        self.assertTrue(Causdocu.objects.filter(pk=causale.pk).exists())
        self.assertTrue(Ord_Mast.objects.filter(pk=documento.pk).exists())
