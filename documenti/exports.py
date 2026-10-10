import io
import json
import re
from decimal import Decimal, InvalidOperation

from django.contrib.auth.decorators import login_required
from django.core.exceptions import RequestDataTooBig
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


MAX_EXPORT_BYTES = 10 * 1024 * 1024
MAX_EXPORT_SHEETS = 20
MAX_EXPORT_ROWS = 50000
MAX_EXPORT_COLUMNS = 100
MAX_CELL_LENGTH = 32767
EXPORT_NAME_PATTERN = re.compile(r'^[A-Za-z0-9_-]{1,80}$')


def _cell_value(value):
    if not isinstance(value, str):
        return value

    testo = value.strip()
    if not testo:
        return None

    valuta = '€' in testo
    numero = testo.replace('€', '').replace('\xa0', '').replace(' ', '')
    if ',' in numero:
        numero = numero.replace('.', '').replace(',', '.')
    elif not valuta and re.fullmatch(r'-?\d{1,3}(?:\.\d{3})+', numero):
        numero = numero.replace('.', '')

    if re.fullmatch(r'-?(?:0|[1-9]\d*)(?:\.\d+)?', numero):
        try:
            decimale = Decimal(numero)
        except InvalidOperation:
            return value
        if decimale == decimale.to_integral_value():
            return int(decimale)
        return decimale

    return value


def _worksheet_title(value, used_titles):
    titolo = re.sub(r'[\[\]:*?/\\]', ' ', value).strip().strip("'")
    titolo = (titolo or 'Elenco')[:31]
    candidato = titolo
    suffisso = 2
    while candidato.casefold() in used_titles:
        finale = f' ({suffisso})'
        candidato = f'{titolo[:31 - len(finale)]}{finale}'
        suffisso += 1
    used_titles.add(candidato.casefold())
    return candidato


@login_required
@require_POST
def esporta_tabelle_excel(request):
    try:
        corpo_richiesta = request.body
    except RequestDataTooBig:
        return JsonResponse(
            {'error': 'Il contenuto da esportare supera il limite consentito.'},
            status=413,
        )
    if len(corpo_richiesta) > MAX_EXPORT_BYTES:
        return JsonResponse(
            {'error': 'Il contenuto da esportare supera il limite consentito.'},
            status=413,
        )

    try:
        payload = json.loads(corpo_richiesta)
    except (json.JSONDecodeError, UnicodeDecodeError):
        return JsonResponse({'error': 'Richiesta di esportazione non valida.'}, status=400)

    tabelle = payload.get('sheets') if isinstance(payload, dict) else None
    nome_pagina = payload.get('page_name') if isinstance(payload, dict) else None
    if (
        not isinstance(tabelle, list)
        or not tabelle
        or len(tabelle) > MAX_EXPORT_SHEETS
    ):
        return JsonResponse(
            {'error': 'Non sono state trovate tabelle da esportare.'},
            status=400,
        )
    if not isinstance(nome_pagina, str) or not EXPORT_NAME_PATTERN.fullmatch(
        nome_pagina
    ):
        return JsonResponse(
            {'error': 'Nome pagina non valido per il file Excel.'},
            status=400,
        )

    workbook = Workbook()
    workbook.remove(workbook.active)
    nomi_usati = set()
    for tabella in tabelle:
        if not isinstance(tabella, dict):
            return JsonResponse({'error': 'Dati tabella non validi.'}, status=400)
        nome = tabella.get('name')
        righe = tabella.get('rows')
        if (
            not isinstance(nome, str)
            or not isinstance(righe, list)
            or not righe
            or len(righe) > MAX_EXPORT_ROWS + 1
        ):
            return JsonResponse({'error': 'Dati tabella non validi.'}, status=400)

        foglio = workbook.create_sheet(_worksheet_title(nome, nomi_usati))
        for indice_riga, riga in enumerate(righe, start=1):
            if not isinstance(riga, list) or len(riga) > MAX_EXPORT_COLUMNS:
                return JsonResponse(
                    {'error': 'Formato delle righe non valido.'},
                    status=400,
                )
            for indice_colonna, valore in enumerate(riga, start=1):
                if not isinstance(valore, (str, int, float, bool, type(None))):
                    return JsonResponse(
                        {'error': 'Valore cella non valido.'},
                        status=400,
                    )
                if isinstance(valore, str) and len(valore) > MAX_CELL_LENGTH:
                    return JsonResponse(
                        {'error': 'Una cella supera la lunghezza massima Excel.'},
                        status=400,
                    )

                cella = foglio.cell(
                    row=indice_riga,
                    column=indice_colonna,
                    value=_cell_value(valore),
                )
                if isinstance(valore, str) and isinstance(cella.value, str):
                    cella.data_type = 's'
                if indice_riga == 1:
                    cella.font = Font(bold=True, color='FFFFFF')
                    cella.fill = PatternFill(
                        fill_type='solid',
                        fgColor='1F4E78',
                    )
                elif isinstance(valore, str) and valore.lstrip().startswith('='):
                    cella.data_type = 's'

        foglio.freeze_panes = 'A2'
        foglio.auto_filter.ref = foglio.dimensions
        for colonna in range(1, foglio.max_column + 1):
            lunghezza_massima = max(
                (
                    len(str(foglio.cell(row=riga, column=colonna).value or ''))
                    for riga in range(1, min(foglio.max_row, 100) + 1)
                ),
                default=10,
            )
            foglio.column_dimensions[get_column_letter(colonna)].width = min(
                max(lunghezza_massima + 2, 10),
                60,
            )

    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        ),
    )
    timestamp = timezone.localtime().strftime('%Y%m%d_%H%M%S')
    response['Content-Disposition'] = (
        f'attachment; filename="{nome_pagina}_{timestamp}.xlsx"'
    )
    return response
